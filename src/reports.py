"""Write CSV, Excel, and PDF one-pager outputs."""

from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from src.allocation import AllocationRow, summarize

PT = ZoneInfo("America/Los_Angeles")


def _period_label(start: datetime, end: datetime) -> str:
    return f"{start.date().isoformat()}_to_{end.date().isoformat()}"


def write_csv(
    rows: list[AllocationRow],
    path: Path,
    period_start: datetime,
    period_end: datetime,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = [
        "period_start",
        "period_end",
        "email",
        "name",
        "category",
        "project",
        "repo_match",
        "spend_usd",
        "ai_lines",
        "allocation_pct",
        "notes",
    ]
    with path.open("w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow(
                {
                    "period_start": period_start.date().isoformat(),
                    "period_end": period_end.date().isoformat(),
                    "email": r.email,
                    "name": r.name or "",
                    "category": r.category,
                    "project": r.project,
                    "repo_match": r.repo_match,
                    "spend_usd": f"{r.spend_usd:.4f}",
                    "ai_lines": r.ai_lines,
                    "allocation_pct": f"{r.allocation_pct:.4f}",
                    "notes": r.notes,
                }
            )
    return path


def write_xlsx(
    rows: list[AllocationRow],
    path: Path,
    period_start: datetime,
    period_end: datetime,
    summary: dict[str, Any] | None = None,
    meta: dict[str, Any] | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    summary = summary or summarize(rows)
    wb = Workbook()

    # Detail
    ws = wb.active
    ws.title = "Detail"
    headers = [
        "period_start",
        "period_end",
        "email",
        "name",
        "category",
        "project",
        "repo_match",
        "spend_usd",
        "ai_lines",
        "allocation_pct",
        "notes",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                period_start.date().isoformat(),
                period_end.date().isoformat(),
                r.email,
                r.name or "",
                r.category,
                r.project,
                r.repo_match,
                r.spend_usd,
                r.ai_lines,
                r.allocation_pct,
                r.notes,
            ]
        )

    # Summary_by_Project
    ws_p = wb.create_sheet("Summary_by_Project")
    ws_p.append(["project", "spend_usd"])
    for cell in ws_p[1]:
        cell.font = Font(bold=True)
    for proj, usd in summary["by_project"].items():
        ws_p.append([proj, usd])

    # Summary_by_Category
    ws_c = wb.create_sheet("Summary_by_Category")
    ws_c.append(["category", "spend_usd"])
    for cell in ws_c[1]:
        cell.font = Font(bold=True)
    for cat in ("CapEx", "OpEx", "Unallocated"):
        ws_c.append([cat, summary["by_category"].get(cat, 0.0)])

    # Coverage
    ws_cov = wb.create_sheet("Coverage")
    ws_cov.append(["metric", "value"])
    for cell in ws_cov[1]:
        cell.font = Font(bold=True)
    ws_cov.append(["total_usd", summary["total_usd"]])
    ws_cov.append(["coverage_pct", summary["coverage_pct"]])
    ws_cov.append(["capex_headcount", summary["capex_headcount"]])
    ws_cov.append(["opex_headcount", summary["opex_headcount"]])
    ws_cov.append(["unallocated_headcount", summary["unallocated_headcount"]])
    if meta:
        for k, v in meta.items():
            ws_cov.append([k, v])

    wb.save(path)
    return path


def write_pdf_one_pager(
    path: Path,
    summary: dict[str, Any],
    period_start: datetime,
    period_end: datetime,
    customer: str = "Cirrus",
    engagement: str = "Tungsten",
    extra_notes: list[str] | None = None,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    generated = datetime.now(PT).strftime("%Y-%m-%d %H:%M %Z")

    doc = SimpleDocTemplate(
        str(path),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom",
        parent=styles["Heading1"],
        fontSize=16,
        spaceAfter=6,
    )
    h2 = ParagraphStyle(
        "H2Custom",
        parent=styles["Heading2"],
        fontSize=12,
        spaceBefore=12,
        spaceAfter=6,
    )
    body = styles["Normal"]

    story = []
    story.append(
        Paragraph(
            f"{customer} / {engagement} — CapEx vs OpEx Chargeback",
            title_style,
        )
    )
    story.append(
        Paragraph(
            f"Period: {period_start.date().isoformat()} → "
            f"{period_end.date().isoformat()} &nbsp;|&nbsp; Generated: {generated}",
            body,
        )
    )
    story.append(Spacer(1, 10))

    cats = summary["by_category"]
    total = summary["total_usd"]
    coverage = summary["coverage_pct"]

    totals_data = [
        ["Category", "Amount (USD)"],
        ["CapEx", f"${cats.get('CapEx', 0):,.2f}"],
        ["OpEx", f"${cats.get('OpEx', 0):,.2f}"],
        ["Unallocated", f"${cats.get('Unallocated', 0):,.2f}"],
        ["Total", f"${total:,.2f}"],
        ["Coverage %", f"{coverage:.1f}%"],
    ]
    t = Table(totals_data, colWidths=[2.5 * inch, 2 * inch])
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTNAME", (0, -2), (-1, -1), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#e8f5e9")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -3), [colors.white, colors.HexColor("#f5f5f5")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(t)
    story.append(
        Paragraph(
            "Coverage % = (CapEx + OpEx) / (CapEx + OpEx + Unallocated). "
            "Unallocated is intentional — never silent CapEx.",
            body,
        )
    )

    story.append(Paragraph("Top projects (by allocated $)", h2))
    proj_rows = [["Project", "Amount (USD)"]]
    for i, (proj, usd) in enumerate(summary["by_project"].items()):
        if i >= 8:
            break
        proj_rows.append([proj, f"${usd:,.2f}"])
    if len(proj_rows) == 1:
        proj_rows.append(["(none)", "$0.00"])
    tp = Table(proj_rows, colWidths=[3 * inch, 2 * inch])
    tp.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16213e")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(tp)

    story.append(Paragraph("Headcount", h2))
    story.append(
        Paragraph(
            f"CapEx people: {summary['capex_headcount']} &nbsp;|&nbsp; "
            f"OpEx people: {summary['opex_headcount']} &nbsp;|&nbsp; "
            f"Unallocated people: {summary['unallocated_headcount']}",
            body,
        )
    )

    if extra_notes:
        story.append(Paragraph("Notes", h2))
        for n in extra_notes:
            story.append(Paragraph(f"• {n}", body))

    doc.build(story)
    return path


def default_output_stems(
    out_dir: Path, period_end: datetime
) -> dict[str, Path]:
    day = period_end.date().isoformat()
    return {
        "csv": out_dir / f"capex_opex_allocation_{day}.csv",
        "xlsx": out_dir / f"capex_opex_allocation_{day}.xlsx",
        "pdf": out_dir / f"one_pager_{day}.pdf",
    }
